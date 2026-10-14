"""Analyst-style respondent data sheets for HubSpot CX survey exports.

LEGACY: client-specific (one client's CX survey). Only used when settings.client_rules_legacy is on;
replaced by the codebook-driven data workbook (IMPLEMENTATION_PLAN.md M7).

Reproduces how a client's CX analysts prepare a HubSpot export by hand (reference:
the analysts' manual workbook): a FULL DATA sheet of respondent rows, with derived 1/blank
flag columns inserted right after the column they come from, then one sheet per
reporting category holding just that category's respondents.

Every sheet has the same layout:
  row 1  headers
  row 2  % of the sheet's respondents (count ÷ base in A3)
  row 3  counts — SUM of each flag column, COUNT of numeric questions, base in A3
  row 4+ respondents

Derived columns:
  COUNT                 1 per respondent (first column)
  Job role              (##) role group flags — one role per respondent, by priority
  Customer type         Number of products, Single/Dual/Triple product, product mix
  NPS question          Promoters / Passives / Detractors
  * Drivers columns     DRIVERS (whichever driver question was answered) + one flag per driver
  1-5 rating questions  one flag column per score (1 … 5)
"""

import io
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from services.data_loader import (
    MULTI_SELECT_DELIMITER, _JOB_ROLE_HEADER, _NPS_HEADER, _job_role_group, _normalize,
)

ROLE_FLAGS = {
    1.0: "(##) Exams officer (regardless of other roles)",
    2.0: "(##) Leader (if not EO, incl HoD who aren't teachers)",
    3.0: "(##) Teachers (All other roles)",
    4.0: "(##) Other (no listed role, e.g. only Other)",
}

# One sheet per reporting category, in the analysts' order: (sheet title, flag column)
CATEGORY_SHEETS = [
    ("Detractors", "Detractors"),
    ("Promoters", "Promoters"),
    ("Single", "Single product"),
    ("Dual", "Dual product"),
    ("Triple", "Triple product"),
    ("Exam officers", ROLE_FLAGS[1.0]),
    ("Teachers", ROLE_FLAGS[3.0]),
    ("Leaders", ROLE_FLAGS[2.0]),
]

# HubSpot export bookkeeping the analysts delete before working on the data
HUBSPOT_SYSTEM_COLUMNS = {
    "response", "industry standard question type", "sentiment", "submission name", "source",
    "survey id", "survey type", "rating", "visitor id", "tags",
}

_CUSTOMER_TYPE_HEADER = re.compile(r"\bcustomer type\b", re.IGNORECASE)
_DRIVERS_HEADER = re.compile(r"\bdrivers\b", re.IGNORECASE)
_PRODUCTS = {  # product → wording it is recognised by in the Customer type answers
    "Resources": "teaching and learning resources",
    "Programmes": "programmes",
    "Exams": "exams or qualifications",
}
# Applied to the exact set of answers given — any other answer (e.g. "Other") rules a mix out
_PRODUCT_MIX = [
    ("Resources only", lambda p: p == {"Resources"}),
    ("Programmes only", lambda p: p == {"Programmes"}),
    ("Exams only", lambda p: p == {"Exams"}),
    ("Resources + Programmes", lambda p: p == {"Resources", "Programmes"}),
    ("Programmes + Exams", lambda p: p == {"Programmes", "Exams"}),
    ("Programmes AND/OR exams", lambda p: bool(p) and p <= {"Programmes", "Exams"}),
    ("Resources + Exams", lambda p: p == {"Resources", "Exams"}),
]

# Driver flag columns in the analysts' order; unknown drivers follow, "Other" last
_DRIVER_ORDER = ["accessibility", "quality", "ease of working", "support", "communication",
                 "outcomes", "value for money"]

HEADER_FONT = Font(bold=True)
TOTALS_FONT = Font(bold=True, color="9A3412")
FLAG_FILL = PatternFill("solid", fgColor="FFF7ED")
PERCENT_FORMAT = "0.0%"

Column = tuple[str, pd.Series, str | None]  # (header, values, totals formula: "SUM" | "COUNT" | None)


def add_respondent_sheets(excel_bytes: bytes, raw: pd.DataFrame, filename: str = "") -> bytes:
    """Append the FULL DATA sheet and one sheet per category to a tabulation workbook."""
    columns = build_columns(raw)
    wb = load_workbook(io.BytesIO(excel_bytes))

    full_title = "FULL DATA PULSE" if "pulse" in filename.lower() else "FULL DATA"
    write_sheet(wb.create_sheet(full_title), columns)

    flags = {header: values for header, values, _ in columns}
    for title, flag in CATEGORY_SHEETS:
        if flag not in flags:
            continue
        rows = flags[flag].eq(1).to_numpy()
        write_sheet(wb.create_sheet(title), [(h, v[rows], t) for h, v, t in columns])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_columns(raw: pd.DataFrame) -> list[Column]:
    """All sheet columns, raw and derived, in sheet order."""
    raw = raw.loc[:, [c for c in raw.columns if _normalize(str(c)) not in HUBSPOT_SYSTEM_COLUMNS]]
    raw = raw.reset_index(drop=True)
    columns: list[Column] = [("COUNT", pd.Series(1.0, index=raw.index), None)]

    headers = [str(h) for h in raw.columns]
    # "Other drivers" is the free-text follow-up, not a driver list
    driver_cols = [h for h, c in zip(headers, raw.columns)
                   if _DRIVERS_HEADER.search(h) and not _normalize(h).startswith("other")
                   and not pd.api.types.is_numeric_dtype(raw[c])]

    for header, col in zip(headers, raw.columns):
        series = raw[col]
        is_rating = _is_rating(series)
        is_text = not pd.api.types.is_numeric_dtype(series)
        columns.append((header, series, "COUNT" if is_rating or _is_nps(header, series) else None))

        if is_text and _JOB_ROLE_HEADER.search(header):
            group = _job_role_group(header, _option_flags(series))
            if group is not None:
                columns += [(title, _flag(group == code), "SUM") for code, title in ROLE_FLAGS.items()]
        if is_text and _CUSTOMER_TYPE_HEADER.search(header):
            columns += _product_columns(series)
        if _is_nps(header, series):
            columns += [
                ("Promoters", _flag(series >= 9), "SUM"),
                ("Passives", _flag(series.between(7, 8)), "SUM"),
                ("Detractors", _flag(series <= 6), "SUM"),
            ]
        if driver_cols and header == driver_cols[-1]:
            columns += _driver_columns(raw, driver_cols)
        if is_rating:
            columns += [(score, _flag(series == score), "SUM") for score in range(1, 6)]
    return columns


def write_sheet(ws, columns: list[Column]) -> None:
    n = len(columns[0][1]) if columns else 0
    first, last = 4, 3 + n

    for j, (header, values, total) in enumerate(columns, start=1):
        letter = get_column_letter(j)
        ws.cell(row=1, column=j, value=header).font = HEADER_FONT
        if j == 1:  # base: number of respondents on this sheet
            ws.cell(row=3, column=1, value=f"=SUM(A{first}:A{last})" if n else 0).font = TOTALS_FONT
        elif total:
            count = ws.cell(row=3, column=j, value=f"={total}({letter}{first}:{letter}{last})" if n else 0)
            count.font = TOTALS_FONT
            share = ws.cell(row=2, column=j, value=f"=IF($A$3=0,0,{letter}3/$A$3)")
            share.font, share.number_format = TOTALS_FONT, PERCENT_FORMAT
        for i, value in enumerate(values, start=first):
            if value is None or (not isinstance(value, str) and pd.isna(value)):
                continue
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime().replace(tzinfo=None)
            cell = ws.cell(row=i, column=j, value=value)
            if total == "SUM":
                cell.fill = FLAG_FILL
    ws.freeze_panes = "C4"


def _selections(series: pd.Series) -> pd.Series:
    """Each answer as a list of the options ticked ("A; B" → ["A", "B"]); [] when blank."""
    return series.map(lambda v: [o.strip() for o in str(v).split(MULTI_SELECT_DELIMITER) if o.strip()]
                      if pd.notna(v) else [])


def _option_flags(series: pd.Series) -> list[tuple[str, pd.Series]]:
    """[(option, 1/0 flags, NaN when unanswered)] — the shape data_loader uses for multi-selects."""
    selections = _selections(series)
    answered = series.notna()
    options = list(dict.fromkeys(o for chosen in selections for o in chosen))
    return [(o, selections.map(lambda s, o=o: 1.0 if o in s else 0.0).where(answered)) for o in options]


def _flag(mask: pd.Series) -> pd.Series:
    """1 where true, blank otherwise — the analysts' flag convention."""
    return mask.fillna(False).astype(bool).map({True: 1, False: None})


def _is_nps(header: str, series: pd.Series) -> bool:
    if not _NPS_HEADER.search(header) or not pd.api.types.is_numeric_dtype(series):
        return False
    values = series.dropna()
    return not values.empty and values.min() >= 0 and values.max() <= 10 and values.max() >= 7


def _is_rating(series: pd.Series) -> bool:
    """Integer 1-5 rating question (e.g. CX principles, desired outcomes)."""
    if not pd.api.types.is_numeric_dtype(series) or pd.api.types.is_bool_dtype(series):
        return False
    values = series.dropna()
    return (values.nunique() >= 3 and values.min() >= 1 and values.max() <= 5
            and bool((values % 1 == 0).all()))


def _product_of(answer: str) -> str | None:
    norm = _normalize(answer)
    return next((name for name, wording in _PRODUCTS.items() if wording in norm), None)


def _product_columns(series: pd.Series) -> list[Column]:
    """Product counts and mix from Customer type, as the analysts derive them."""
    items = _selections(series)
    # Matches the analysts' =LEN(M3)-LEN(SUBSTITUTE(M3,";",""))+1 — every listed answer counts
    n_products = items.map(lambda v: float(len(v)) if v else None)
    # Exact answer set; a non-product answer is kept as itself so it rules out every mix
    answers = items.map(lambda v: {_product_of(a) or a for a in v})

    columns: list[Column] = [
        ("Number of products", n_products, "COUNT"),
        ("Single product", _flag(n_products == 1), "SUM"),
        ("Dual product", _flag(n_products == 2), "SUM"),
        ("Triple product", _flag(n_products >= 3), "SUM"),
    ]
    for title, rule in _PRODUCT_MIX:
        columns.append((title, _flag(answers.map(lambda p, r=rule: r(p))), "SUM"))
    return columns


def _driver_order(driver: str) -> tuple:
    norm = _normalize(driver)
    if norm == "other":
        return (2, 0, norm)
    known = next((i for i, key in enumerate(_DRIVER_ORDER) if key in norm), None)
    return (0, known, norm) if known is not None else (1, 0, norm)


def _driver_columns(raw: pd.DataFrame, driver_cols: list[str]) -> list[Column]:
    """DRIVERS = the driver answer given (detractor, passive or promoter), then a flag per driver."""
    by_header = {str(c): raw[c] for c in raw.columns}
    combined = pd.Series(None, index=raw.index, dtype=object)
    for header in driver_cols:
        combined = combined.where(combined.notna(), by_header[header])

    selections = _selections(combined)
    drivers: list[str] = []
    for chosen in selections:
        drivers += [d for d in chosen if d not in drivers]
    drivers.sort(key=_driver_order)

    columns: list[Column] = [("DRIVERS", combined, None)]
    for driver in drivers:
        columns.append((driver, _flag(selections.map(lambda s, d=driver: d in s)), "SUM"))
    return columns
