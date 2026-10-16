"""Data workbook: respondent-level sheets built from the codebook (RECODING_PLAN.md §7, R6).

A FULL DATA sheet of every respondent with the file's own answers, plus 1/blank flag columns
inserted next to the column they come from, then one sheet per reporting category holding just
that category's respondents. What goes in is a `DataWorkbook` setting on the codebook (set in the
review page or by a setup), so any survey can get the layout analysts build by hand.

Every sheet has the same layout:
  row 1  headers
  row 2  % of the sheet's respondents (count ÷ base) — unless `percent_row` is off
  row 3  counts — SUM of each flag column, COUNT of numeric questions, base in column A
  row 4+ respondents                                 (rows 2 / 3+ without the % row)

Columns:
  COUNT                       1 per respondent (first column)
  every column of the file    as answered, except `drop_columns`
  `groups`                    extra columns after a given column: flags for one answer of a
                              variable (`var` + `code`) or a condition (`when`), one flag per
                              answer / option (`flags_of`), number of options ticked (`count`),
                              or the first answered of several columns (`first_of`)
  `score_flags` [lo, hi]      after each rating question with scores in lo..hi: one flag per score
"""

import io
import re
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from services.codebook.apply import apply_codebook
from services.codebook.model import Codebook, DataWorkbook, SheetColumn
from services.codebook.rules import Context, RuleError, evaluate
from services.ingest.readers import read_raw

HEADER_FONT = Font(bold=True)
TOTALS_FONT = Font(bold=True, color="9A3412")
FLAG_FILL = PatternFill("solid", fgColor="FFF7ED")
PERCENT_FORMAT = "0.0%"
MAX_SHEET_TITLE = 31

Column = tuple[Any, pd.Series, str | None]  # (header, values, totals formula: "SUM" | "COUNT" | None)


def _norm(text: Any) -> str:
    return re.sub(r"\s+", " ", str(text)).strip().casefold()


def _flag(mask: pd.Series) -> pd.Series:
    """1 where true, blank otherwise — the analysts' flag convention."""
    return mask.fillna(False).astype(bool).map({True: 1, False: None})


class _Builder:
    """Resolves column specs against one file: its raw rows, codebook and coded variables."""

    def __init__(self, df: pd.DataFrame, codebook: Codebook):
        self.df, self.cb = df, codebook
        coded = apply_codebook(df, codebook)
        self.columns = {k: v.reset_index(drop=True) for k, v in coded.columns.items()}
        self.ctx = Context(codebook, self.columns, df.index)
        self.warnings: list[str] = []

    def raw(self, header: str) -> pd.Series | None:
        match = next((c for c in self.df.columns if _norm(c) == _norm(header)), None)
        return None if match is None else self.df[match]

    def coded(self, ref: str) -> pd.Series:
        var = self.cb.find(ref)
        name = var.name if var is not None and var.name else ref
        if name not in self.columns:
            raise RuleError(f"'{ref}' is not a coded variable")
        return self.columns[name]

    def options(self, ref: str) -> list[tuple[str, pd.Series]]:
        """(label, 0/1 series) per answer of a multi-select / multi-output derived, or per code."""
        if self.ctx.is_multi(ref):
            return list(self.ctx.options(ref).items())
        var, d = self.cb.find(ref), self.cb.find_derived(ref)
        codes = var.codes if var is not None else (d.codes if d is not None else [])
        if not codes:
            raise RuleError(f"'{ref}' has no answers to flag")
        series = self.coded(ref)
        return [(c.label, series.eq(c.code).astype(float).where(series.notna())) for c in codes]

    def expand(self, spec: SheetColumn) -> list[Column]:
        if spec.flags_of:
            items = self.options(spec.flags_of)
            items = _ordered(items, spec.order, spec.last)
            headers = {_norm(k): v for k, v in spec.headers.items()}
            out = []
            for label, values in items:
                code = _code_of(self.cb, spec.flags_of, label)
                header = headers.get(_norm(label)) or (headers.get(_norm(f"{code:g}")) if code is not None else None) or label
                out.append((header, _flag(values.eq(1)), "SUM"))
            return out
        if spec.count:
            skip = {_norm(x) for x in spec.exclude}
            ticked = pd.concat([v.eq(1) for label, v in self.ctx.options(spec.count).items()
                                if _norm(label) not in skip], axis=1)
            n = ticked.sum(axis=1).astype(float).where(self.ctx.answered(spec.count))
            return [(spec.header, n.map(lambda v: None if pd.isna(v) else float(v)), spec.total or "COUNT")]
        if spec.first_of:
            combined = pd.Series(None, index=self.df.index, dtype=object)
            for header in spec.first_of:
                series = self.raw(header)
                if series is not None:
                    combined = combined.where(combined.notna(), series)
            return [(spec.header, combined, spec.total)]
        if spec.when is not None:
            return [(spec.header, _flag(evaluate(spec.when, self.ctx)), spec.total or "SUM")]
        if spec.var is not None and spec.code is not None:
            return [(spec.header, _flag(self.coded(spec.var).eq(spec.code)), spec.total or "SUM")]
        if spec.var is not None:
            return [(spec.header or spec.var, self.coded(spec.var), spec.total)]
        raise RuleError(f"column '{spec.header}' says nothing to show")

    def safe(self, spec: SheetColumn) -> list[Column]:
        try:
            return self.expand(spec)
        except (RuleError, KeyError, ValueError) as e:
            self.warnings.append(f"Left out column {spec.header or spec.flags_of or ''}: {e}")
            return []

    def rows_of(self, tab) -> pd.Series | None:
        try:
            if tab.when is not None:
                return evaluate(tab.when, self.ctx).fillna(False).astype(bool)
            return self.coded(tab.var).eq(tab.code).fillna(False)
        except (RuleError, KeyError, ValueError) as e:
            self.warnings.append(f"Left out tab {tab.title}: {e}")
            return None


def _code_of(cb: Codebook, ref: str, label: str) -> float | None:
    var, d = cb.find(ref), cb.find_derived(ref)
    codes = var.codes if var is not None and var.role != "multi" else (d.codes if d is not None else [])
    return next((c.code for c in codes if c.label == label), None)


def _ordered(items: list[tuple[str, pd.Series]], order: list[str], last: list[str]) -> list[tuple[str, pd.Series]]:
    """Keyword order (`order`: substrings), then the rest alphabetically, then `last` (exact answers)."""
    if not order and not last:
        return items
    keys, tail = [_norm(k) for k in order], [_norm(k) for k in last]

    def key(item):
        norm = _norm(item[0])
        if norm in tail:
            return (2, tail.index(norm), norm)
        known = next((i for i, k in enumerate(keys) if k in norm), None)
        return (0, known, norm) if known is not None else (1, 0, norm)

    return sorted(items, key=key)


def _is_scored(var, lo: int, hi: int) -> bool:
    if var is None or var.role != "numeric" or len(var.codes) < 3:
        return False
    return all(float(c.code).is_integer() and lo <= c.code <= hi for c in var.codes)


def build_columns(df: pd.DataFrame, codebook: Codebook, settings: DataWorkbook) -> tuple[list[Column], _Builder]:
    """All FULL DATA columns in sheet order (the builder holds the warnings)."""
    b = _Builder(df, codebook)
    dropped = {_norm(c) for c in settings.drop_columns}
    by_column = {_norm(v.column): v for v in codebook.variables}
    after: dict[str, list[SheetColumn]] = {}
    names = {}
    for group in settings.groups:
        after.setdefault(_norm(group.after), []).extend(group.columns)
        names[_norm(group.after)] = group.after
    placed = set()

    columns: list[Column] = [("COUNT", pd.Series(1.0, index=df.index), None)]
    for col in df.columns:
        header, norm = str(col), _norm(col)
        if norm in dropped:
            continue
        var = by_column.get(norm)
        numeric_question = var is not None and var.role == "numeric" and bool(var.codes)
        columns.append((header, df[col], "COUNT" if numeric_question else None))
        for spec in after.get(norm, []):
            columns += b.safe(spec)
        placed.add(norm)
        if settings.score_flags and _is_scored(var, *settings.score_flags):
            lo, hi = settings.score_flags
            columns += [(score, _flag(df[col] == score), "SUM") for score in range(lo, hi + 1)]
    for key in after:
        if key not in placed:
            b.warnings.append(f"Column '{names[key]}' isn't in this file: the columns after it were left out")
    return columns, b


def write_sheet(ws, columns: list[Column], percent_row: bool = True) -> None:
    n = len(columns[0][1]) if columns else 0
    counts = 3 if percent_row else 2
    first, last = counts + 1, counts + n

    for j, (header, values, total) in enumerate(columns, start=1):
        letter = get_column_letter(j)
        ws.cell(row=1, column=j, value=header).font = HEADER_FONT
        if j == 1:  # base: number of respondents on this sheet
            ws.cell(row=counts, column=1, value=f"=SUM(A{first}:A{last})" if n else 0).font = TOTALS_FONT
        elif total:
            count = ws.cell(row=counts, column=j, value=f"={total}({letter}{first}:{letter}{last})" if n else 0)
            count.font = TOTALS_FONT
            if percent_row:
                share = ws.cell(row=2, column=j, value=f"=IF($A$3=0,0,{letter}3/$A$3)")
                share.font, share.number_format = TOTALS_FONT, PERCENT_FORMAT
        for i, value in enumerate(values, start=first):
            if value is None or (not isinstance(value, str) and pd.isna(value)):
                continue
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime().replace(tzinfo=None)
            cell = ws.cell(row=i, column=j, value=value)
            if total == "SUM":
                cell.fill = FLAG_FILL
    ws.freeze_panes = f"C{first}"


def _title(wb, title: str) -> str:
    title = re.sub(r"[\[\]:*?/\\]", " ", title).strip()[:MAX_SHEET_TITLE] or "Sheet"
    taken = {ws.title.casefold() for ws in wb.worksheets}
    candidate, i = title, 2
    while candidate.casefold() in taken:
        suffix = f" ({i})"
        candidate, i = title[:MAX_SHEET_TITLE - len(suffix)] + suffix, i + 1
    return candidate


def add_sheets(wb, df: pd.DataFrame, codebook: Codebook, settings: DataWorkbook) -> list[str]:
    """Append FULL DATA + one sheet per tab to an openpyxl workbook → warnings."""
    df = df.dropna(axis=0, how="all").reset_index(drop=True)
    columns, b = build_columns(df, codebook, settings)
    write_sheet(wb.create_sheet(_title(wb, settings.title)), columns, settings.percent_row)
    for tab in settings.tabs:
        rows = b.rows_of(tab)
        if rows is None:
            continue
        mask = rows.to_numpy()
        write_sheet(wb.create_sheet(_title(wb, tab.title)), [(h, v[mask], t) for h, v, t in columns],
                    settings.percent_row)
    return b.warnings


def _frame(file_bytes: bytes, filename: str) -> pd.DataFrame:
    return read_raw(file_bytes, filename).df.dropna(axis=0, how="all").reset_index(drop=True)


def append_data_workbook(excel_bytes: bytes, file_bytes: bytes, filename: str, codebook: Codebook) -> bytes:
    """Blocking: the tables workbook with the codebook's data sheets appended (unchanged if none set)."""
    if codebook is None or codebook.data_workbook is None:
        return excel_bytes
    wb = load_workbook(io.BytesIO(excel_bytes))
    add_sheets(wb, _frame(file_bytes, filename), codebook, codebook.data_workbook)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def data_workbook_bytes(file_bytes: bytes, filename: str, codebook: Codebook,
                        settings: DataWorkbook | None = None) -> tuple[bytes, list[str]]:
    """Blocking: a workbook with just the data sheets → (xlsx bytes, warnings)."""
    settings = settings or codebook.data_workbook or DataWorkbook()
    wb = Workbook()
    wb.remove(wb.active)
    warnings = add_sheets(wb, _frame(file_bytes, filename), codebook, settings)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), warnings
