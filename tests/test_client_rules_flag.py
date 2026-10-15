"""Client rules are a setup, not inference; CLIENT_RULES_LEGACY now only gates the analyst-style
data sheets (respondent_sheet.py), until the data workbook replaces them (RECODING_PLAN.md R6)."""

import io

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from config import get_settings
from services.data_loader import load_data

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _survey_xlsx() -> bytes:
    rng = np.random.default_rng(3)
    n = 80
    df = pd.DataFrame({
        "Record ID": range(1000, 1000 + n),
        "Job role": rng.choice(["Teacher", "Exams Officer; Teacher", "Head of Department", "School Leader",
                                "Head of Department; Teacher", "Exams Officer"], n),
        "Overall NPS": rng.integers(0, 11, n),
        "Satisfaction": rng.choice(["Agree", "Disagree", "Strongly agree"], n),
    })
    buf = io.BytesIO()
    df.to_excel(buf, index=False, sheet_name="Data")
    return buf.getvalue()


@pytest.fixture
def legacy_off(monkeypatch):
    monkeypatch.setattr(get_settings(), "client_rules_legacy", False)


def test_inference_has_no_client_rules():
    data = load_data(_survey_xlsx(), "s.xlsx")
    assert "Job_role_group" not in data.df.columns  # comes from the HubSpot CX setup instead
    assert "Job_role_Teacher" in data.df.columns  # generic multi-select split still happens


def _auto_analyze_sheets(client, auth_headers) -> list[str]:
    response = client.post(
        "/v1/auto-analyze", headers=auth_headers,
        files={"file": ("survey.xlsx", _survey_xlsx(), XLSX_MIME)},
        data={"options": '{"include_summary": false}'},
    )
    assert response.status_code == 200, response.text
    return load_workbook(io.BytesIO(response.content), read_only=True).sheetnames


def test_auto_analyze_adds_data_sheets_by_default(client, auth_headers):
    assert "FULL DATA" in _auto_analyze_sheets(client, auth_headers)


def test_auto_analyze_skips_data_sheets_without_legacy_rules(client, auth_headers, legacy_off):
    assert not any(s.startswith("FULL DATA") for s in _auto_analyze_sheets(client, auth_headers))
