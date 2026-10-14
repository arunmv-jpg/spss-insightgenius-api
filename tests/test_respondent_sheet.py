"""Analyst-style FULL DATA + per-category sheets for HubSpot CX exports."""

import io

import pandas as pd
from openpyxl import Workbook, load_workbook

from services.respondent_sheet import add_respondent_sheets, build_columns

RES = "Teaching and learning resources such as textbooks, digital tools, or curriculum support"


def _raw():
    return pd.DataFrame({
        "Record ID": [11, 12, 13, 14],
        "Source": ["Shareable Link"] * 4,
        "Job role": ["Teacher; Exams Officer", "Head of Department", "Teacher", "Other"],
        "Customer type": [f"{RES}; Exams or qualifications", "Exams or qualifications",
                          f"Programmes or curriculum; Exams or qualifications; {RES}; Other", f"{RES}; Other"],
        "Overall NPS": [10, 8, 3, 9],
        "Detractors Drivers": [None, None, "The support Acme provide", None],
        "Passive Drivers": [None, "The ease of working with Acme; Other", None, None],
        "Promoter Drivers": ["The quality of Acme's products and services", None, None,
                             "The support Acme provide; The ease of working with Acme"],
        "CX Principle - Ease": [5, 4, 1, 3],
    })


def _cols():
    """{header: values}, blanks as None (pandas holds them as NaN)."""
    return {header: [None if not isinstance(v, str) and pd.isna(v) else v for v in values]
            for header, values, _ in build_columns(_raw())}


def test_columns_follow_their_source_in_analyst_order():
    headers = [h for h, _, _ in build_columns(_raw())]
    assert headers[0] == "COUNT" and "Source" not in headers  # HubSpot bookkeeping dropped
    assert headers[headers.index("Job role") + 1].startswith("(##) Exams officer")
    assert headers[headers.index("Customer type") + 1] == "Number of products"
    assert headers[headers.index("Overall NPS") + 1:headers.index("Overall NPS") + 4] == [
        "Promoters", "Passives", "Detractors"]
    i = headers.index("DRIVERS")
    assert headers[i - 1] == "Promoter Drivers"
    assert headers[i + 1:i + 5] == ["The quality of Acme's products and services",
                                     "The ease of working with Acme", "The support Acme provide", "Other"]
    assert headers[-5:] == [1, 2, 3, 4, 5]


def test_role_flags_one_role_per_respondent():
    cols = _cols()
    assert cols["(##) Exams officer (regardless of other roles)"] == [1, None, None, None]
    assert cols["(##) Leader (if not EO, incl HoD who aren't teachers)"] == [None, 1, None, None]
    assert cols["(##) Teachers (All other roles)"] == [None, None, 1, None]
    assert cols["(##) Other (no listed role, e.g. only Other)"] == [None, None, None, 1]


def test_product_and_nps_flags():
    cols = _cols()
    assert cols["Number of products"] == [2.0, 1.0, 4.0, 2.0]
    assert cols["Dual product"] == [1, None, None, 1]
    assert cols["Triple product"] == [None, None, 1, None]  # 3 products + Other counts as triple
    assert cols["Resources + Exams"] == [1, None, None, None]
    assert cols["Exams only"] == [None, 1, None, None]
    assert cols["Resources only"] == [None, None, None, None]  # "Other" alongside rules a mix out
    assert cols["Programmes AND/OR exams"] == [None, 1, None, None]
    assert cols["Promoters"] == [1, None, None, 1]
    assert cols["Detractors"] == [None, None, 1, None]


def test_drivers_combined_and_flagged():
    cols = _cols()
    assert cols["DRIVERS"][2] == "The support Acme provide"
    assert cols["The support Acme provide"] == [None, None, 1, 1]
    assert cols["Other"] == [None, 1, None, None]
    assert cols[5] == [1, None, None, None]


def test_full_data_and_category_sheets():
    wb = Workbook()
    wb.active.title = "Summary"
    buf = io.BytesIO()
    wb.save(buf)

    wb = load_workbook(io.BytesIO(add_respondent_sheets(buf.getvalue(), _raw(), "cx-pulse_w3.xlsx")))
    assert wb.sheetnames == ["Summary", "FULL DATA PULSE", "Detractors", "Promoters", "Single", "Dual",
                             "Triple", "Exam officers", "Teachers", "Leaders"]

    full = wb["FULL DATA PULSE"]
    headers = [c.value for c in full[1]]
    col = headers.index("Promoters") + 1
    letter = full.cell(row=1, column=col).column_letter
    assert full["A3"].value == "=SUM(A4:A7)"
    assert full.cell(row=3, column=col).value == f"=SUM({letter}4:{letter}7)"
    assert full.cell(row=2, column=col).value == f"=IF($A$3=0,0,{letter}3/$A$3)"
    assert full.cell(row=3, column=headers.index("Overall NPS") + 1).value.startswith("=COUNT(")
    assert full.max_row == 7 and full.freeze_panes == "C4"

    # Category sheets hold only their respondents, with their own totals
    promoters = wb["Promoters"]
    assert [promoters.cell(row=r, column=2).value for r in range(4, promoters.max_row + 1)] == [11, 14]
    assert promoters["A3"].value == "=SUM(A4:A5)"
    assert [c.value for c in wb["Teachers"]["B"][3:]] == [13]
