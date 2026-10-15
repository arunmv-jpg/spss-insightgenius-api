"""Data workbook: FULL DATA + category sheets built from the codebook (RECODING_PLAN.md R6)."""

import io
import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from services.codebook.infer import infer_codebook
from services.codebook.patch import EditError, apply_ops
from services.data_loader import read_raw
from services.data_workbook import build_columns, data_workbook_bytes
from services.setups import apply_setup, get_setup

FIXTURE = Path(__file__).parent / "fixtures" / "exports" / "hubspot_cx_sample.xlsx"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
RES = "Teaching and learning resources such as textbooks, digital tools, or curriculum support"

# Four respondents with known answers on top of the synthetic HubSpot export
KNOWN = {
    "Record ID": [11, 12, 13, 14],
    "Job role": ["Teacher; Exams Officer", "Head of Department", "Teacher", "Other"],
    "Customer type": [f"{RES}; Exams or qualifications", "Exams or qualifications",
                      f"Programmes or curriculum; Exams or qualifications; {RES}; Other", f"{RES}; Other"],
    "CX Pulse Overall NPS": [10, 8, 3, 9],
    "Detractors Drivers": [None, None, "The support Acme provide", None],
    "Passive Drivers": [None, "The ease of working with Acme; Other", None, None],
    "Promoter Drivers": ["The quality of Acme's products and services", None, None,
                         "The support Acme provide; The ease of working with Acme"],
    "CX Principle - Ease": [5, 4, 1, 3],
}


@pytest.fixture(autouse=True)
def _fresh_rate_limit():
    from middleware import rate_limiter
    yield
    rate_limiter._windows.clear()


def _xlsx(df: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    return buf.getvalue()


@pytest.fixture(scope="module")
def hubspot():
    df = pd.read_excel(FIXTURE)
    for col, values in KNOWN.items():
        df[col] = df[col].astype(object)
        df.loc[:3, col] = values
    file_bytes = _xlsx(df)
    raw = read_raw(file_bytes, "cx-pulse_w3.xlsx")
    cb, _, skipped = apply_setup(infer_codebook(raw), get_setup("hubspot_cx_pulse"), raw.df)
    assert not skipped
    return file_bytes, raw, cb


def _cols(hubspot):
    _, raw, cb = hubspot
    columns, b = build_columns(raw.df.dropna(how="all").reset_index(drop=True), cb, cb.data_workbook)
    assert not b.warnings
    first4 = lambda vals: [None if not isinstance(v, str) and pd.isna(v) else v for v in vals[:4]]
    cols = {h: first4(vals) for h, vals, _ in columns}
    ease = next(i for i, (h, _, _) in enumerate(columns) if h == "CX Principle - Ease")
    cols.update({("Ease", h): first4(vals) for h, vals, _ in columns[ease + 1:ease + 6]})  # score headers repeat
    return [h for h, _, _ in columns], cols


def test_columns_follow_their_source_in_analyst_order(hubspot):
    headers, _ = _cols(hubspot)
    assert headers[0] == "COUNT" and "Source" not in headers  # HubSpot bookkeeping dropped
    assert headers[headers.index("Job role") + 1].startswith("(##) Exams officer")
    assert headers[headers.index("Customer type") + 1] == "Number of products"
    nps = headers.index("CX Pulse Overall NPS")
    assert headers[nps + 1:nps + 4] == ["Promoters", "Passives", "Detractors"]
    i = headers.index("DRIVERS")
    assert headers[i - 1] == "Promoter Drivers"
    drivers = headers[i + 1:headers.index("Other drivers")]  # the free-text follow-up comes next
    assert drivers[0] == "The accessibility of Acme's products and services" and drivers[-1] == "Other"
    assert drivers.index("The quality of Acme's products and services") < drivers.index("The ease of working with Acme")
    ease = headers.index("CX Principle - Ease")
    assert headers[ease + 1:ease + 6] == [1, 2, 3, 4, 5]


def test_role_product_nps_driver_and_score_flags(hubspot):
    _, cols = _cols(hubspot)
    assert cols["(##) Exams officer (regardless of other roles)"] == [1, None, None, None]
    assert cols["(##) Leader (if not EO, incl HoD who aren't teachers)"] == [None, 1, None, None]
    assert cols["(##) Teachers (All other roles)"] == [None, None, 1, None]
    assert cols["(##) Other (no listed role, e.g. only Other)"] == [None, None, None, 1]
    assert cols["Number of products"] == [2.0, 1.0, 4.0, 2.0]
    assert cols["Dual product"] == [1, None, None, 1]
    assert cols["Triple product"] == [None, None, 1, None]  # 3 products + Other counts as triple
    assert cols["Resources + Exams"] == [1, None, None, None]
    assert cols["Exams only"] == [None, 1, None, None]
    assert cols["Resources only"] == [None, None, None, None]  # "Other" alongside rules a mix out
    assert cols["Programmes AND/OR exams"] == [None, 1, None, None]
    assert cols["Promoters"] == [1, None, None, 1]
    assert cols["Detractors"] == [None, None, 1, None]
    assert cols["DRIVERS"][2] == "The support Acme provide"
    assert cols["The support Acme provide"] == [None, None, 1, 1]
    assert cols["Other"] == [None, 1, None, None]
    assert cols[("Ease", 5)] == [1, None, None, None] and cols[("Ease", 1)] == [None, None, 1, None]


def test_full_data_and_category_sheets(hubspot):
    file_bytes, raw, cb = hubspot
    xlsx, warnings = data_workbook_bytes(file_bytes, "cx-pulse_w3.xlsx", cb)
    assert warnings == []
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["FULL DATA PULSE", "Detractors", "Promoters", "Single", "Dual",
                             "Triple", "Exam officers", "Teachers", "Leaders"]
    full = wb["FULL DATA PULSE"]
    headers = [c.value for c in full[1]]
    col = headers.index("Promoters") + 1
    letter = full.cell(row=1, column=col).column_letter
    last = 3 + len(raw.df)
    assert full["A3"].value == f"=SUM(A4:A{last})"
    assert full.cell(row=3, column=col).value == f"=SUM({letter}4:{letter}{last})"
    assert full.cell(row=2, column=col).value == f"=IF($A$3=0,0,{letter}3/$A$3)"
    assert full.cell(row=3, column=headers.index("CX Pulse Overall NPS") + 1).value.startswith("=COUNT(")
    assert full.freeze_panes == "C4"
    # Category sheets hold only their respondents
    record = headers.index("Record ID") + 1
    promoters = wb["Promoters"]
    assert [promoters.cell(row=r, column=record).value for r in (4, 5)] == [11, 14]
    nps = raw.df["CX Pulse Overall NPS"]
    assert promoters.max_row - 3 == int((nps >= 9).sum())
    assert wb["Teachers"].cell(row=4, column=record).value == 13


# ── any survey ───────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def survey():
    rng = np.random.default_rng(7)
    n = 90
    df = pd.DataFrame({
        "Email": [f"p{i}@example.com" for i in range(n)],
        "Region": rng.choice(["North", "South", "East"], n),
        "Satisfaction": rng.integers(1, 6, n),
        "Brands used": rng.choice(["Alpha; Beta", "Beta", "Alpha; Beta; Gamma", "Gamma"], n),
    })
    file_bytes = _xlsx(df)
    raw = read_raw(file_bytes, "survey.xlsx")
    return file_bytes, raw, infer_codebook(raw), df


SETTINGS = {
    "title": "Respondents", "drop_columns": ["Email"], "score_flags": [1, 5],
    "groups": [
        {"after": "Region", "columns": [{"flags_of": "Region"}]},
        {"after": "Brands used", "columns": [
            {"header": "Brands", "count": "Brands used"},
            {"header": "Uses Gamma", "when": {"var": "Brands used", "any": ["Gamma"]}}]},
        {"after": "Not a column", "columns": [{"header": "x", "var": "Region", "code": 1}]},
    ],
    "tabs": [{"title": "North", "var": "Region", "code": 2}, {"title": "Ghost", "var": "Nope", "code": 1}],
}


def test_any_survey_gets_its_own_sheets(survey):
    file_bytes, raw, cb, df = survey
    cb = apply_ops(cb, [{"op": "set_data_workbook", "settings": SETTINGS}], raw.df)
    region = cb.variable("Region")
    assert {c.label: c.code for c in region.codes}["North"] == 2  # natural order: East, North, South
    xlsx, warnings = data_workbook_bytes(file_bytes, "survey.xlsx", cb)
    assert any("Not a column" in w for w in warnings) and any("Ghost" in w for w in warnings)
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["Respondents", "North"]
    rows = list(wb["Respondents"].iter_rows(values_only=True))
    headers = list(rows[0])
    assert "Email" not in headers
    assert headers[headers.index("Region") + 1:headers.index("Region") + 4] == ["East", "North", "South"]
    assert headers[headers.index("Brands used") + 1:headers.index("Brands used") + 3] == ["Brands", "Uses Gamma"]
    sat = headers.index("Satisfaction")
    assert headers[sat + 1:sat + 6] == [1, 2, 3, 4, 5]  # a 1-5 rating: one flag per score after it
    data = rows[3:]
    brands = [r[headers.index("Brands")] for r in data]
    assert brands == [float(len(v.split(";"))) for v in df["Brands used"]]
    assert wb["North"].max_row - 3 == int((df["Region"] == "North").sum())


def test_settings_validated_and_cleared(survey):
    _, raw, cb, _ = survey
    with pytest.raises(EditError, match="needs a variable and a category"):
        apply_ops(cb, [{"op": "set_data_workbook", "settings": {"tabs": [{"title": "x"}]}}], raw.df)
    with pytest.raises(EditError, match="invalid data workbook"):
        apply_ops(cb, [{"op": "set_data_workbook", "settings": {"groups": [{"columns": []}]}}], raw.df)
    on = apply_ops(cb, [{"op": "set_data_workbook", "settings": SETTINGS}], raw.df)
    assert apply_ops(on, [{"op": "set_data_workbook", "settings": None}], raw.df).data_workbook is None


# ── API ──────────────────────────────────────────────────────────────────────

def _post(client, auth_headers, path, data, file_bytes, name):
    return client.post(path, headers=auth_headers, data=data, files={"file": (name, file_bytes, XLSX_MIME)})


def test_export_endpoint(client, auth_headers, survey):
    file_bytes, raw, cb, _ = survey
    r = _post(client, auth_headers, "/v1/exports/data-workbook",
              {"codebook": cb.model_dump_json(), "settings": json.dumps(SETTINGS)}, file_bytes, "survey.xlsx")
    assert r.status_code == 200, r.text
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["Respondents", "North"]
    assert "Ghost" in r.headers["X-Warnings"]
    # No settings anywhere: just the FULL DATA sheet
    r = _post(client, auth_headers, "/v1/exports/data-workbook", {}, file_bytes, "survey.xlsx")
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["FULL DATA"]


def test_auto_analyze_appends_the_sheets_only_when_set(client, auth_headers, hubspot):
    file_bytes, _, cb = hubspot
    opts = {"options": '{"include_summary": false}'}
    r = _post(client, auth_headers, "/v1/auto-analyze", {**opts, "codebook": cb.model_dump_json()}, file_bytes, "x.xlsx")
    assert r.status_code == 200, r.text
    assert "Leaders" in load_workbook(io.BytesIO(r.content), read_only=True).sheetnames
    r = _post(client, auth_headers, "/v1/auto-analyze", opts, file_bytes, "x.xlsx")
    assert not any(s.startswith("FULL DATA") for s in load_workbook(io.BytesIO(r.content), read_only=True).sheetnames)
