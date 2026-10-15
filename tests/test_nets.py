"""Nets (Top 2 Box etc.) in the codebook, and answer values in means (RECODING_PLAN.md R5)."""

import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from services.codebook.infer import infer_codebook
from services.codebook.patch import EditError, apply_ops
from services.ingest.readers import read_raw

FIXTURE = Path(__file__).parent / "fixtures" / "exports" / "hubspot_cx_sample.xlsx"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXAM = "Exam_Results_Alternative"  # 1 Strongly Disagree … 5 Strongly Agree; 98/99 excluded
T2B = {"label": "Agree (T2B)", "codes": [4, 5]}


@pytest.fixture(autouse=True)
def _fresh_rate_limit():
    from middleware import rate_limiter
    yield
    rate_limiter._windows.clear()


@pytest.fixture(scope="module")
def raw():
    return read_raw(FIXTURE.read_bytes(), FIXTURE.name)


@pytest.fixture(scope="module")
def cb(raw):
    return infer_codebook(raw)


def nets_of(cb, var=EXAM):
    return [(n.label, n.codes) for n in cb.variable(var).nets]


def test_set_add_remove_nets(raw, cb):
    edited = apply_ops(cb, [{"op": "set_nets", "var": EXAM, "nets": [T2B, {"label": "Disagree (B2B)", "codes": ["Strongly Disagree", "Disagree"]}]}], raw.df)
    assert nets_of(edited) == [("Agree (T2B)", [4.0, 5.0]), ("Disagree (B2B)", [1.0, 2.0])]
    edited = apply_ops(edited, [{"op": "remove_net", "var": EXAM, "label": "agree (t2b)"},
                                {"op": "add_net", "var": EXAM, "label": "Neutral", "codes": [3]}], raw.df)
    assert nets_of(edited) == [("Disagree (B2B)", [1.0, 2.0]), ("Neutral", [3.0])]
    assert edited.table_nets() == {EXAM: {"Disagree (B2B)": [1, 2], "Neutral": [3]}}
    assert edited.table_nets({EXAM: {"x": [1]}}) == {EXAM: {"x": [1]}}  # a request's own nets win


@pytest.mark.parametrize("net, message", [
    ({"label": "", "codes": [4]}, "needs a label"),
    ({"label": "X", "codes": []}, "at least one answer"),
    ({"label": "X", "codes": [98]}, "excluded from the base"),
    ({"label": "X", "codes": [42]}, "not a code"),
])
def test_invalid_nets_rejected(raw, cb, net, message):
    with pytest.raises(EditError, match=message):
        apply_ops(cb, [{"op": "add_net", "var": EXAM, **net}], raw.df)


def test_duplicate_net_names_rejected(raw, cb):
    with pytest.raises(EditError, match="two nets"):
        apply_ops(cb, [{"op": "set_nets", "var": EXAM, "nets": [T2B, {"label": "agree (T2B)", "codes": [1]}]}], raw.df)


def test_nets_follow_code_changes(raw, cb):
    base = apply_ops(cb, [{"op": "set_nets", "var": EXAM, "nets": [T2B, {"label": "Bottom", "codes": [1]}]}], raw.df)
    reversed_ = apply_ops(base, [{"op": "set_codes", "var": EXAM, "map": [[1, 5], [2, 4], [4, 2], [5, 1]]}], raw.df)
    assert nets_of(reversed_) == [("Agree (T2B)", [2.0, 1.0]), ("Bottom", [5.0])]  # same answers, new values
    merged = apply_ops(base, [{"op": "merge_codes", "var": EXAM, "codes": [5], "into": 4}], raw.df)
    assert nets_of(merged)[0] == ("Agree (T2B)", [4.0])
    excluded = apply_ops(base, [{"op": "set_missing", "var": EXAM, "code": 1, "missing": True}], raw.df)
    assert nets_of(excluded) == [("Agree (T2B)", [4.0, 5.0])]  # "Bottom" lost its only answer
    retyped = apply_ops(base, [{"op": "set_role", "var": EXAM, "role": "text"}], raw.df)
    assert nets_of(retyped) == []


# ── tables ───────────────────────────────────────────────────────────────────

def _post(client, auth_headers, path, data):
    return client.post(path, headers=auth_headers, data=data,
                       files={"file": (FIXTURE.name, FIXTURE.read_bytes(), XLSX_MIME)})


def _rows(excel: bytes, first_cell: str) -> list[tuple]:
    wb = load_workbook(io.BytesIO(excel))
    return [r for ws in wb.worksheets for r in ws.iter_rows(values_only=True) if r and str(r[0]).strip() == first_cell]


def test_tabulate_shows_codebook_nets_and_means_use_values(client, auth_headers, raw, cb):
    spec = json.dumps({"banner": "CX_Pulse_Overall_NPS_group", "stubs": [EXAM], "include_means": True})
    plain = apply_ops(cb, [{"op": "add_net", "var": EXAM, **T2B}], raw.df)
    r = _post(client, auth_headers, "/v1/tabulate", {"spec": spec, "codebook": plain.model_dump_json()})
    assert r.status_code == 200, r.text
    answers = raw.df["Exam Results Alternative"]
    expected = 100 * answers.isin(["Agree", "Strongly Agree"]).sum() / answers.isin(
        ["Strongly Disagree", "Disagree", "Neither agree or disagree", "Agree", "Strongly Agree"]).sum()
    net_row = _rows(r.content, "Agree (T2B)")[0]
    assert float(str(net_row[1]).rstrip("%")) == pytest.approx(expected, abs=0.1)
    mean = float(str(_rows(r.content, "Mean")[0][1]).split()[0])

    # Reverse the scale (values 5…1): the net still covers Agree + Strongly Agree, the mean flips
    flipped = apply_ops(plain, [{"op": "set_codes", "var": EXAM, "map": [[1, 5], [2, 4], [4, 2], [5, 1]]}], raw.df)
    r = _post(client, auth_headers, "/v1/tabulate", {"spec": spec, "codebook": flipped.model_dump_json()})
    assert r.status_code == 200, r.text
    assert _rows(r.content, "Agree (T2B)")[0][1] == net_row[1]
    assert float(str(_rows(r.content, "Mean")[0][1]).split()[0]) == pytest.approx(6 - mean, abs=0.011)


def test_auto_analyze_uses_codebook_nets(client, auth_headers, raw, cb):
    edited = apply_ops(cb, [{"op": "add_net", "var": EXAM, "label": "Any disagree", "codes": [1, 2]}], raw.df)
    r = _post(client, auth_headers, "/v1/auto-analyze", {"codebook": edited.model_dump_json(),
                                                          "options": '{"include_summary": false}'})
    assert r.status_code == 200, r.text
    assert _rows(r.content, "Any disagree")
