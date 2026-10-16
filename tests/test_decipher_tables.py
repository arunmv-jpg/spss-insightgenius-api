"""Decipher-style tables: two significance levels (CAPITALS / lowercase) + low-base markers."""

import io
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from services.data_loader import load_data
from services.quantipy_engine import QuantiProEngine
from services.tabulation_builder import TabulateSpec, _group_letters, _sig_suffix, build_tabulation
from shared.significance import low_base_marker, significant_at

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(df: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    df.to_excel(buf, index=False, sheet_name="Data")
    return buf.getvalue()


def _survey() -> pd.DataFrame:
    """Group A says Yes 100/200 times, B 78/200: p ≈ 0.027. With 3 columns Bonferroni divides
    alpha by 3 pairs (95% → 0.0167, 90% → 0.0333), so A>B is significant at 90% but not 95%.
    Group C has only 20 respondents (low base)."""
    rows = ([("A", "Yes")] * 100 + [("A", "No")] * 100
            + [("B", "Yes")] * 78 + [("B", "No")] * 122
            + [("C", "Yes")] * 10 + [("C", "No")] * 10)
    df = pd.DataFrame(rows, columns=["Group", "Answer"])
    df["Score"] = [5 if a == "Yes" else 3 for a in df["Answer"]]
    return df


@pytest.fixture(scope="module")
def data():
    return load_data(_xlsx(_survey()), "survey.xlsx")


# ── helpers ──────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("base,marker", [(29, "**"), (30, "*"), (99, "*"), (100, ""), (5000, "")])
def test_low_base_marker_boundaries(base, marker):
    assert low_base_marker(base) == marker


def test_low_base_marker_can_be_switched_off():
    assert low_base_marker(10, warn=0, flag=0) == ""


def test_significant_at_needs_higher_column_and_level():
    assert significant_at(0.07, 0.10, 2, higher=True)
    assert not significant_at(0.07, 0.10, 2, higher=False)  # only the higher column gets a letter
    assert not significant_at(0.07, None, 2, higher=True)   # no secondary level requested
    assert not significant_at(0.07, 0.10, 3, higher=True)   # Bonferroni: 0.10 / 3 pairs


def test_sig_suffix_and_positional_group_letters():
    class Col:
        def __init__(self, letter):
            self.letter = letter
    group = [Col("E"), Col("F"), Col("G")]
    assert _group_letters(["A", "C"], group) == ["E", "G"]
    assert _group_letters(["B"], [Col("B"), Col("C")]) == ["C"]  # positional, not by name
    assert _sig_suffix(["E"], ["G"]) == "Eg"


# ── engine ───────────────────────────────────────────────────────────────────

def _cell(result, answer_label, group_value):
    row = next(r for r in result["table"] if r["row_label"] == answer_label)
    return row[group_value]


def test_engine_secondary_letters(data):
    groups = {v: k for k, v in data.meta.variable_value_labels["Group"].items()}
    a = str(groups["A"])
    single = QuantiProEngine.crosstab_with_significance(data, "Answer", "Group", significance_level=0.95)
    two = QuantiProEngine.crosstab_with_significance(data, "Answer", "Group", significance_level=0.95,
                                                     secondary_level=0.90)
    b_letter = two["col_labels"][str(groups["B"])]

    assert "significance_letters_secondary" not in _cell(single, "Yes", a)  # unchanged by default
    assert b_letter not in _cell(two, "Yes", a)["significance_letters"]
    assert b_letter in _cell(two, "Yes", a)["significance_letters_secondary"]
    assert two["significance_level_secondary"] == 0.90


# ── workbook ─────────────────────────────────────────────────────────────────

def _workbook(data, **spec_kwargs):
    spec = TabulateSpec(banners=["Group"], stubs=["Answer", "Score"], include_means=True, **spec_kwargs)
    return load_workbook(io.BytesIO(build_tabulation(QuantiProEngine, data, spec).excel_bytes))


def _row(ws, label):
    return next(r for r in ws.iter_rows(values_only=True) if r[0] == label)


def test_workbook_lowercase_letters_and_legend(data):
    ws = _workbook(data, significance_levels=[0.90, 0.95])["Answer"]  # order doesn't matter
    yes = _row(ws, "Yes")
    assert yes[2] == "50.0% b"  # column A beats B at 90% only
    assert "95% (CAPITALS) / 90% (lowercase)" in ws["A2"].value
    assert "** base < 30" in ws["A2"].value


def test_workbook_single_level_has_no_lowercase(data):
    ws = _workbook(data, significance_level=0.95)["Answer"]
    assert not any(isinstance(v, str) and v.split()[-1].islower() for v in _row(ws, "Yes")[1:] if v)


def test_workbook_low_base_number_format(data):
    ws = _workbook(data)["Answer"]
    base_row = next(r for r in ws.iter_rows() if r[0].value == "Base (N)")
    by_value = {c.value: c.number_format for c in base_row[1:]}
    assert by_value[20] == '0" **"'      # group C, n=20
    assert by_value[200] == "General"    # groups A/B, n=200: no marker
    assert isinstance(base_row[1].value, int)  # still a number


def test_summary_sheet_rows(data):
    summary = {r[0]: r[1] for r in _workbook(data, significance_levels=[0.95, 0.90])["Summary"]
               .iter_rows(values_only=True) if r[0]}
    assert summary["Significance Level"] == "95% (CAPITALS) / 90% (lowercase)"
    assert summary["Low Base Markers"] == "** base < 30, * base < 100"


def test_invalid_levels_rejected():
    with pytest.raises(ValueError):
        TabulateSpec(significance_levels=[95, 90])


# ── API ──────────────────────────────────────────────────────────────────────

@pytest.fixture(autouse=True)
def _fresh_rate_limit():
    """Don't spend the shared test key's per-minute budget other test modules rely on."""
    from middleware import rate_limiter
    yield
    rate_limiter._windows.clear()


def test_crosstab_endpoint_levels_and_low_base(client, auth_headers):
    spec = {"row": "Answer", "col": "Group", "significance_levels": [0.95, 0.90]}
    response = client.post("/v1/crosstab", headers=auth_headers,
                           files={"file": ("s.xlsx", _xlsx(_survey()), XLSX_MIME)},
                           data={"spec": json.dumps(spec)})
    assert response.status_code == 200, response.text
    body = response.json()["data"]
    assert body["significance_level_secondary"] == 0.90
    assert sorted(body["column_bases"].values()) == [20, 200, 200]
    assert list(body["low_base"].values()) == ["**"]


def test_tabulate_endpoint_rejects_bad_levels(client, auth_headers):
    spec = {"banners": ["Group"], "stubs": ["Answer"], "significance_levels": [95]}
    response = client.post("/v1/tabulate", headers=auth_headers,
                           files={"file": ("s.xlsx", _xlsx(_survey()), XLSX_MIME)},
                           data={"spec": json.dumps(spec)})
    assert response.status_code == 400
